"""
error_direction.py

Computes error DIRECTION (systematic over- vs. under-estimation), not just
magnitude, for every approach this module can get real per-timestep
predictions for — across every real-life dataset x trim combination.


Metric: NMBE (Normalized Mean Bias Error) = mean(pred - actual) / mean(|actual|),
computed per (dataset, trim, series, approach[, regime]) — positive = systematic
overestimation, negative = underestimation.

Sanity check: every row also carries the MAE computed from these same raw
predictions, plus (when a matching entry exists) the MAE independently recorded
in results/ for that exact (approach, regime, dataset, trim)

"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ts_comparison import (
    ROOT, RESULTS, BEST, GLSTM, STAT_FORECASTERS,
    load_splits, load_foundation_predictions, get_recorded_foundation_mae,
    load_ppm_raw_predictions, _ppm_metrics_path,
)
from split_correlation import (
    REAL_TRIM_NAMES, real_logs, NAIVE_BASELINES, STATISTICAL,
    DISPLAY_NAMES as _SC_DISPLAY_NAMES, parse_approach_maes, parse_ppm_maes,
    _val_mean_mae,
)
from extra_forecasters import EXTRA_FORECASTERS

synth_logs = [
    p.stem for p in sorted((ROOT / "data" / "synthetic").glob("*.xes"))
    if "recency" not in p.stem
]

STAT_MODELS = ["naive", "ets", "sarimax", "theta", "stl"]  # prophet handled separately (own HPO source)
ML_MODELS = ["ridge", "ridge_mimo", "gru", "gru_mimo"]
DARTS_MODELS = ["nbeats", "tft"]
FOUNDATION_MODELS = ["chronos", "tabpfn"]
SIMULATION_MODELS = ["pmsd", "simod", "agentsimulator"]
PPM_MODELS = ["camargo", "amiri", "bukhsh_rt", "bukhsh_suffix"]
PPM_REGIMES = ["plain_field", "full", "half"]  # display order: plain -> full -> half

APPROACH_ORDER = (["naive", "val_mean"] + STAT_MODELS[1:] + ["prophet"] + ML_MODELS
                  + DARTS_MODELS + FOUNDATION_MODELS + SIMULATION_MODELS + PPM_MODELS)
DISPLAY_NAMES = {**_SC_DISPLAY_NAMES, "gru": "GRU", "gru_mimo": "GRU_mimo",
                 "val_mean": "Val. Avg.", "amiri": "PGT", "nbeats": "N-BEATS",
                 "simod": "Simod", "agentsimulator": "AgentSimulator"}
REGIME_DISPLAY = {"plain_field": "Plain", "full": "First", "half": "Half"}
SERIES = ["concurrent_cases", "throughput_time"]
TS_PREDICTIONS_ROOT = ROOT / "results" / "ts_predictions"


def _nmbe(pred: np.ndarray, actual: np.ndarray) -> float:
    m = np.mean(np.abs(actual))
    return float(np.mean(pred - actual) / m) if m != 0 else np.nan


def _pct_over(pred: np.ndarray, actual: np.ndarray) -> float:
    return float(np.mean(pred > actual))


def _row(dataset, trim, approach, regime, series, pred, actual, recorded_mae=None) -> dict:
    scale = float(np.mean(np.abs(actual)))
    mae = float(np.mean(np.abs(pred - actual)))
    mae_match = None
    if recorded_mae is not None:
        mae_match = abs(mae - recorded_mae) < max(1e-6, 0.01 * abs(recorded_mae))
    return {
        "dataset": dataset, "trim": trim, "approach": approach, "regime": regime, "series": series,
        "nmbe": _nmbe(pred, actual), "pct_over": _pct_over(pred, actual), "n": len(actual),
        "mae": mae, "recorded_mae": recorded_mae, "mae_match": mae_match,
        "actual_scale": scale,
        "unstable_scale": scale < 5,
    }



def _pick_time_csv(dir_path: Path):
    """Same logic as training_times.py's _pick_time_csv (kept independent —
    this module has no dependency on training_times.py)."""
    candidates = [c for c in sorted(dir_path.glob("time*.csv"))
                  if not c.name.startswith("time_prophet_")]
    if not candidates:
        return None
    named = [c for c in candidates if c.name != "time.csv"]
    pool = named or candidates
    return max(pool, key=lambda p: p.stat().st_mtime)


def _best_params_for(model: str, dataset: str, trim: str, is_real: bool, series: str) -> dict:
    """Best recorded hyperparams for a TS model, from its own tune/fit
    history — time_{dataset}.csv for naive/ets/sarimax/stl/ridge/ridge_mimo/gru/
    gru_mimo, time_prophet_{dataset}.csv for prophet (the real HPO grid added
    in prophet_baseline.py)."""
    d = (RESULTS / trim / dataset) if is_real else (RESULTS / "synthetic" / "none" / dataset)
    time_path = (d / f"time_prophet_{dataset}.csv") if model == "prophet" else _pick_time_csv(d)
    if time_path is None or not time_path.exists():
        return {}
    tdf = pd.read_csv(time_path)
    row = tdf[(tdf["model"] == model) & (tdf["series"] == series) & (tdf["phase"] == "fit")]
    if row.empty:
        return {}
    val = row.iloc[0]["params_json"]
    if pd.isna(val):
        return {}
    import json
    return json.loads(val)


def _ts_predictions_path(model: str, dataset: str, trim: str, is_real: bool) -> Path:
    sub = trim if is_real else "synthetic/none"
    return TS_PREDICTIONS_ROOT / model / sub / dataset / "predictions.csv"


def _read_cached_ts_predictions(model: str, dataset: str, trim: str, is_real: bool) -> dict:
    """Returns {series: {date_str: value}} for whatever's already persisted at
    _ts_predictions_path, or {} if nothing cached yet for this combo."""
    p = _ts_predictions_path(model, dataset, trim, is_real)
    if not p.exists():
        return {}
    df = pd.read_csv(p)
    out = {}
    for s in SERIES:
        sub = df[df["series"] == s]
        if not sub.empty:
            out[s] = dict(zip(sub["date"].astype(str), sub["predicted"]))
    return out


def _write_ts_prediction_series(model: str, dataset: str, trim: str, is_real: bool,
                                series_name: str, pred_by_date: dict) -> None:
    """Persists one series' predictions, merging with (never clobbering)
    whatever's already saved there for the OTHER series."""
    p = _ts_predictions_path(model, dataset, trim, is_real)
    p.parent.mkdir(parents=True, exist_ok=True)
    new_rows = pd.DataFrame([{"series": series_name, "date": d, "predicted": v}
                             for d, v in pred_by_date.items()])
    if p.exists():
        existing = pd.read_csv(p)
        existing = existing[existing["series"] != series_name]
        if not existing.empty:
            new_rows = pd.concat([existing, new_rows], ignore_index=True)
    new_rows.to_csv(p, index=False)


def _recorded_ts_mae(model: str, dataset: str, trim: str, is_real: bool, series_name: str):
    """Recorded MAE straight from metrics_<dataset>.csv — used only as a
    sanity check against the (cached or freshly-computed) prediction series'
    own MAE, the same file/row naive/ets/.../prophet already share for their
    magnitude-only comparisons elsewhere (mae_comparison_table, results_auto.xlsx)."""
    d = (RESULTS / trim / dataset) if is_real else (RESULTS / "synthetic" / "none" / dataset)
    p = d / f"metrics_{dataset}.csv"
    if not p.exists():
        return None
    df = pd.read_csv(p)
    row = df[(df["model"] == model) & (df["series"] == series_name)]
    return float(row.iloc[0]["mae"]) if not row.empty else None


def _get_or_compute_ts_prediction(model: str, dataset: str, trim: str, is_real: bool,
                                  series_name: str, test_idx, compute_fn) -> np.ndarray | None:
    """compute_fn: () -> array-like of predictions aligned to test_idx, called
    ONLY if nothing is cached yet for this exact (model, dataset, trim, series)."""
    cached = _read_cached_ts_predictions(model, dataset, trim, is_real).get(series_name)
    if cached and all(str(d.date()) in cached for d in test_idx):
        return np.array([cached[str(d.date())] for d in test_idx], dtype=float)
    try:
        pred = np.asarray(compute_fn(), dtype=float)
    except Exception as e:
        print(f"  [warn] {dataset}/{trim} {model}/{series_name}: refit failed: {e}")
        return None
    pred_by_date = {str(d.date()): float(v) for d, v in zip(test_idx, pred)}
    _write_ts_prediction_series(model, dataset, trim, is_real, series_name, pred_by_date)
    return pred


def get_ts_model_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    from prophet_baseline import _fit_prophet as _prophet_fit

    try:
        from time_series_prediction import forecast_nbeats, forecast_tft
        _darts_forecasters = {"nbeats": forecast_nbeats, "tft": forecast_tft}
    except Exception as e:
        _darts_forecasters = {}
        print(f"  [warn] darts forecasters unavailable ({e}) -- skipping nbeats/tft")

    rows = []
    for series_name in SERIES:
        s = split["cc"] if series_name == "concurrent_cases" else split["tt"]
        train_val = pd.concat([s["train"], s["val"]])
        test_idx = s["test"].index
        actual = s["test"].to_numpy()

        val_avg = float(s["val"].mean())
        pred = _get_or_compute_ts_prediction(
            "val_mean", dataset, trim, is_real, series_name, test_idx,
            lambda v=val_avg, n=len(test_idx): np.full(n, v))
        if pred is not None:
            rec = _val_mean_mae(dataset, trim, is_real).get(series_name)
            rows.append(_row(dataset, trim, "val_mean", None, series_name, pred, actual, recorded_mae=rec))

        for model in STAT_MODELS + ML_MODELS:
            forecaster = STAT_FORECASTERS.get(model) or EXTRA_FORECASTERS.get(model)
            params = _best_params_for(model, dataset, trim, is_real, series_name)
            pred = _get_or_compute_ts_prediction(
                model, dataset, trim, is_real, series_name, test_idx,
                lambda fc=forecaster, p=params: fc(train_val, len(test_idx), p))
            if pred is None:
                continue
            rec = _recorded_ts_mae(model, dataset, trim, is_real, series_name)
            rows.append(_row(dataset, trim, model, None, series_name, pred, actual, recorded_mae=rec))

        for model in DARTS_MODELS:
            forecaster = _darts_forecasters.get(model)
            if forecaster is None:
                continue
            params = _best_params_for(model, dataset, trim, is_real, series_name)
            pred = _get_or_compute_ts_prediction(
                model, dataset, trim, is_real, series_name, test_idx,
                lambda fc=forecaster, p=params: fc(train_val, len(test_idx), p))
            if pred is None:
                continue
            rec = _recorded_ts_mae(model, dataset, trim, is_real, series_name)
            rows.append(_row(dataset, trim, model, None, series_name, pred, actual, recorded_mae=rec))

        prophet_params = _best_params_for("prophet", dataset, trim, is_real, series_name)
        pred = _get_or_compute_ts_prediction(
            "prophet", dataset, trim, is_real, series_name, test_idx,
            lambda p=prophet_params: (_prophet_fit(train_val, test_idx, p)
                                      .reindex(test_idx).ffill().bfill().fillna(0).to_numpy()))
        if pred is not None:
            rec = _recorded_ts_mae("prophet", dataset, trim, is_real, series_name)
            rows.append(_row(dataset, trim, "prophet", None, series_name, pred, actual, recorded_mae=rec))
    return rows


# ── Foundation models (Chronos/TabPFN) ───

def get_foundation_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    foundation = load_foundation_predictions(dataset, trim, is_real)
    recorded = get_recorded_foundation_mae(dataset, trim, is_real)
    rows = []
    for model in FOUNDATION_MODELS:
        if model not in foundation:
            continue
        for series_name in SERIES:
            s = split["cc"] if series_name == "concurrent_cases" else split["tt"]
            preds_by_date = foundation[model].get(series_name, {})
            actual_vals, pred_vals = [], []
            for d in s["test"].index:
                v = preds_by_date.get(str(d.date()))
                if v is not None:
                    actual_vals.append(s["test"][d])
                    pred_vals.append(v)
            if actual_vals:
                rec = recorded.get(model, {}).get(series_name)
                rows.append(_row(dataset, trim, model, None, series_name,
                                 np.array(pred_vals), np.array(actual_vals), recorded_mae=rec))
    return rows


# ── PMSD (simulation baseline)

def _recorded_pmsd_mae(dataset: str, trim: str, is_real: bool, series_name: str):
    """Recorded MAE from PMSD's own metrics_<dataset>_pmsd.csv — deliberately
    a separate file from metrics_<dataset>.csv (see PMSD README), so this needs
    its own reader rather than _recorded_ts_mae."""
    d = (RESULTS / trim / dataset) if is_real else (RESULTS / "synthetic" / "none" / dataset)
    p = d / f"metrics_{dataset}_pmsd.csv"
    if not p.exists():
        return None
    df = pd.read_csv(p)
    row = df[df["series"] == series_name]
    return float(row.iloc[0]["mae"]) if not row.empty else None


def get_pmsd_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    cached = _read_cached_ts_predictions("pmsd", dataset, trim, is_real)
    rows = []
    for series_name in SERIES:
        preds_by_date = cached.get(series_name, {})
        if not preds_by_date:
            continue
        s = split["cc"] if series_name == "concurrent_cases" else split["tt"]
        actual_vals, pred_vals = [], []
        for d in s["test"].index:
            v = preds_by_date.get(str(d.date()))
            if v is not None:
                actual_vals.append(s["test"][d])
                pred_vals.append(v)
        if actual_vals:
            rec = _recorded_pmsd_mae(dataset, trim, is_real, series_name)
            rows.append(_row(dataset, trim, "pmsd", None, series_name,
                             np.array(pred_vals), np.array(actual_vals), recorded_mae=rec))
    return rows


# ── Simod / AgentSimulator (simulation baselines)

def _simulator_predictions_path(model: str, dataset: str) -> Path:
    return RESULTS / model / dataset / f"predictions_{dataset}.csv"


def _recorded_simulator_mae(model: str, dataset: str, series_name: str, trim: str, is_real: bool):
    """Recorded MAE from results/<model>/metrics_<model>[_real].csv (eval_simod.py's
    own output) """
    suffix = "_real" if is_real else ""
    p = RESULTS / model / f"metrics_{model}{suffix}.csv"
    if not p.exists():
        return None
    df = pd.read_csv(p)
    mask = (df["dataset"] == dataset) & (df["series"] == series_name)
    if is_real and "trim" in df.columns:
        mask &= (df["trim"] == trim)
    row = df[mask]
    return float(row.iloc[0]["mae"]) if not row.empty else None


def get_simulator_bias_rows(model: str, dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    """Synthetic (simod, agentsimulator) and real-life ssd (simod only )."""
    p = _simulator_predictions_path(model, dataset)
    if not p.exists():
        return []
    df = pd.read_csv(p)
    if is_real and "trim" in df.columns:
        df = df[df["trim"] == trim]
    rows = []
    for series_name in SERIES:
        sub = df[df["series"] == series_name]
        if sub.empty:
            continue
        preds_by_date = dict(zip(sub["date"].astype(str), sub["predicted"]))
        s = split["cc"] if series_name == "concurrent_cases" else split["tt"]
        actual_vals, pred_vals = [], []
        for d in s["test"].index:
            v = preds_by_date.get(str(d.date()))
            if v is not None:
                actual_vals.append(s["test"][d])
                pred_vals.append(v)
        if actual_vals:
            rec = _recorded_simulator_mae(model, dataset, series_name, trim, is_real)
            rows.append(_row(dataset, trim, model, None, series_name,
                             np.array(pred_vals), np.array(actual_vals), recorded_mae=rec))
    return rows


def get_simod_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    return get_simulator_bias_rows("simod", dataset, trim, is_real, split)


def get_agentsimulator_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    return get_simulator_bias_rows("agentsimulator", dataset, trim, is_real, split)


# ── PPM approaches (full trace + half prefix + plain-field) — raw saved ────


def _recorded_ppm_mae(model: str, regime: str, dataset: str, trim: str, is_real: bool) -> dict:
    """Recorded MAE for a PPM model/regime, read straight from results/ — used only
    as a sanity check against the MAE computed from the raw predictions loaded by
    load_ppm_raw_predictions, never as a prediction source itself.

    """
    run = f"{dataset}_test_full"
    if regime == "plain_field":
        p, _ = _ppm_metrics_path(dataset, trim, model)
        if p is None:
            return {}
        df = pd.read_csv(p)
        return {r["series"]: r["mae"] for _, r in df.iterrows()}

    base_model = "bukhsh" if model in ("bukhsh_rt", "bukhsh_suffix") else model
    subdir = f"{base_model}_hpo_half" if regime == "half" else f"{base_model}_hpo"
    fname = f"metrics_{run}_rt.csv" if model == "bukhsh_rt" else f"metrics_{run}.csv"
    p = RESULTS / subdir / trim / run / fname
    if not p.exists():
        return {}
    df = pd.read_csv(p)
    return {r["series"]: r["mae"] for _, r in df.iterrows()}


def get_ppm_bias_rows(dataset: str, trim: str, is_real: bool, split: dict) -> list[dict]:
    from time_series_creation import (
        create_concurrent_cases_timeseries, create_avg_throughtput_time_timeseries,
    )
    rows = []
    for regime in PPM_REGIMES:
        for model in PPM_MODELS:
            try:
                pred_log = load_ppm_raw_predictions(model, regime, dataset, trim, is_real)
            except Exception as e:
                print(f"  [warn] {dataset}/{trim} {model}/{regime}: raw prediction read failed: {e}")
                pred_log = None
            if pred_log is None or pred_log.empty:
                continue

            pred_cc = create_concurrent_cases_timeseries(
                pred_log, time_col="end_timestamp", case_col="caseid", window="days", plot=False)
            pred_tt = create_avg_throughtput_time_timeseries(
                pred_log, time_col="end_timestamp", case_col="caseid", window="days", plot=False)
            recorded = _recorded_ppm_mae(model, regime, dataset, trim, is_real)

            for series_name, pred_series, actual_series in [
                ("concurrent_cases", pred_cc, split["cc"]["test"]),
                ("throughput_time", pred_tt, split["tt"]["test"]),
            ]:
                pred = pred_series.reindex(actual_series.index).ffill().bfill().fillna(0).to_numpy()
                actual = actual_series.to_numpy()
                rows.append(_row(dataset, trim, model, regime, series_name, pred, actual,
                                 recorded_mae=recorded.get(series_name)))
    return rows


# ── Orchestration ────────────────────────────────────────────────────────────

def build_bias_table(datasets: list[str] | None = None,
                     trims: list[str] | None = None,
                     is_real: bool = True,
                     verbose: bool = True) -> pd.DataFrame:
    """is_real=False builds over synthetic logs instead (datasets defaults to
    synth_logs, trims to ['none'] — synthetic only ever has the 'none' trim)."""
    if is_real:
        datasets = datasets if datasets is not None else real_logs
        trims = trims if trims is not None else REAL_TRIM_NAMES
    else:
        datasets = datasets if datasets is not None else synth_logs
        trims = trims if trims is not None else ["none"]

    rows = []
    for ds in datasets:
        for trim in trims:
            if verbose:
                print(f"=== {ds} / {trim} ===")
            try:
                split = load_splits(ds, trim, is_real)
            except Exception as e:
                if verbose:
                    print(f"  split ERROR: {e}")
                continue
            rows.extend(get_ts_model_bias_rows(ds, trim, is_real, split))
            rows.extend(get_foundation_bias_rows(ds, trim, is_real, split))
            rows.extend(get_pmsd_bias_rows(ds, trim, is_real, split))
            rows.extend(get_simod_bias_rows(ds, trim, is_real, split))
            rows.extend(get_agentsimulator_bias_rows(ds, trim, is_real, split))
            rows.extend(get_ppm_bias_rows(ds, trim, is_real, split))
    return pd.DataFrame(rows)


def nmbe_display_order() -> list[str]:

    return [f"{DISPLAY_NAMES.get(a, a)}" + (f" ({REGIME_DISPLAY[r]})" if a in PPM_MODELS else "")
           for a in APPROACH_ORDER for r in (PPM_REGIMES if a in PPM_MODELS else [None])]


def summarize_bias(df: pd.DataFrame) -> pd.DataFrame:

    key = df.apply(lambda r: f"{DISPLAY_NAMES.get(r['approach'], r['approach'])}"
                             + (f" ({REGIME_DISPLAY[r['regime']]})" if pd.notna(r["regime"]) else ""),
                   axis=1)
    out = df.assign(_key=key).groupby("_key").agg(
        mean_nmbe=("nmbe", "mean"), median_nmbe=("nmbe", "median"),
        mean_pct_over=("pct_over", "mean"),
        n_rows=("nmbe", "count"), n_unstable=("unstable_scale", "sum"),
    )
    out.index.name = "approach"
    order = [o for o in nmbe_display_order() if o in out.index]
    return out.reindex(order)


# ── All-approaches MAE comparison (mirrors compile_results.ipynb) ──────────


AGG_ML_MODELS = ["ridge", "ridge_mimo", "nbeats", "tft"]
MAE_METHOD_ORDER = NAIVE_BASELINES + STATISTICAL + AGG_ML_MODELS + FOUNDATION_MODELS
MAE_DISPLAY_NAMES = {**DISPLAY_NAMES, "nbeats": "N-BEATS"}


MAE_CANONICAL_ORDER = [
    ("Naive", None), ("Val. Avg.", None), ("Prophet", None), ("ETS", None),
    ("SARIMAX", None), ("STL", None), ("Ridge", None), ("Ridge_M", None),
    ("N-BEATS", None), ("TFT", None), ("Chronos", None), ("TabPFN", None),
    ("GLSTM", "plain_field"), ("PGT", "plain_field"), ("PT_RT", "plain_field"), ("PT_su", "plain_field"),
    ("GLSTM", "full"), ("PGT", "full"), ("PT_su", "full"), ("PT_RT", "full"),
    ("GLSTM", "half"), ("PGT", "half"), ("PT_su", "half"), ("PT_RT", "half"),
]


def _nbeats_mae(dataset: str, trim: str, is_real: bool) -> dict:

    d = (RESULTS / trim / dataset) if is_real else (RESULTS / "synthetic" / "none" / dataset)
    p = d / f"metrics_{dataset}.csv"
    if not p.exists():
        return {}
    df = pd.read_csv(p)
    sub = df[df["model"] == "nbeats"]
    return dict(zip(sub["series"], sub["mae"])) if not sub.empty else {}


def mae_comparison_table(datasets: list[str], trim: str, is_real: bool) -> pd.DataFrame:

    sums: dict[tuple, dict] = {}

    def _add(key, series_vals):
        bucket = sums.setdefault(key, {s: [] for s in SERIES})
        for s in SERIES:
            v = series_vals.get(s)
            if v is not None:
                bucket[s].append(v)

    for ds in datasets:
        non_ppm = parse_approach_maes(ds, trim, is_real)
        non_ppm["nbeats"] = _nbeats_mae(ds, trim, is_real)
        for model in MAE_METHOD_ORDER:
            if model in non_ppm:
                _add((MAE_DISPLAY_NAMES.get(model, model), None), non_ppm[model])
        for regime in PPM_REGIMES:
            ppm = parse_ppm_maes(ds, trim, regime)
            for model in PPM_MODELS:
                if model in ppm:
                    _add((MAE_DISPLAY_NAMES.get(model, model), regime), ppm[model])

    rows = []
    for (approach, regime), d in sums.items():
        row = {"approach": approach, "regime": regime}
        for s in SERIES:
            row[f"{s}_mae"] = float(np.mean(d[s])) if d[s] else None
            row[f"{s}_n"] = len(d[s])
        rows.append(row)
    df = pd.DataFrame(rows)

    order_index = {key: i for i, key in enumerate(MAE_CANONICAL_ORDER)}
    df["_ord"] = df.apply(lambda r: order_index.get((r["approach"], r["regime"]), len(MAE_CANONICAL_ORDER)), axis=1)
    return df.sort_values("_ord").drop(columns=["_ord"]).reset_index(drop=True)



_RESULTS_AUTO_ROW = {
    "naive": "naive", "val_mean": "val_average",
    "prophet": "prophet", "ets": "ets", "sarimax": "sarimax", "stl": "stl",
    "ridge": "ridge", "ridge_mimo": "ridge_mimo", "nbeats": "nbeats", "tft": "tft",
    "chronos": "chronos", "tabpfn": "tabpfn",
    "pmsd": "PMSD", "simod": "Simod", "agentsimulator": "AgentSimulator",
}
_RESULTS_AUTO_PPM_ROW = {
    ("camargo", "plain_field"): "PF Camargo", ("amiri", "plain_field"): "PF Amiri",
    ("bukhsh_rt", "plain_field"): "PF Bukhsh RT", ("bukhsh_suffix", "plain_field"): "PF Bukhsh suffix",
    ("camargo", "full"): "GLSTM full", ("amiri", "full"): "Amiri full",
    ("bukhsh_suffix", "full"): "PT suffix full", ("bukhsh_rt", "full"): "PT RT full",
    ("camargo", "half"): "GLSTM half prefix", ("amiri", "half"): "Amiri half",
    ("bukhsh_suffix", "half"): "PT suffix half prefix", ("bukhsh_rt", "half"): "PT RT half prefix",
}

_RA_SYNTH_LOAN = ["loan_flat", "loan_seasonal", "loan_trend", "loan_drift", "loan_combined"]
_RA_SYNTH_O2C  = ["o2c_flat", "o2c_seasonal", "o2c_trend", "o2c_drift", "o2c_combined"]
_RA_REAL_LIFE  = ["bpic12-a", "bpic15-1", "bpic15-2", "bpic17-o",
                  "bpic20-dom", "bpic20-int", "helpdesk", "sepsis"]


def _read_results_auto_mae(trim: str) -> dict:

    import openpyxl
    wb = openpyxl.load_workbook(RESULTS / "results_auto.xlsx", data_only=True)
    ws = wb[trim]
    is_none = trim == "none"
    real_base_col = 40 if is_none else 2

    out: dict[tuple, dict] = {}
    for series, rows in (("concurrent_cases", range(5, 34)), ("throughput_time", range(40, 69))):
        for r in rows:
            model = ws.cell(row=r, column=1).value
            if not model:
                continue
            groups = ([(_RA_SYNTH_LOAN, 2), (_RA_SYNTH_O2C, 18), (_RA_REAL_LIFE, real_base_col)]
                     if is_none else [(_RA_REAL_LIFE, real_base_col)])
            for ds_list, base_col in groups:
                for i, ds in enumerate(ds_list):
                    v = ws.cell(row=r, column=base_col + 2 + i * 3).value
                    out.setdefault((model, ds), {})[series] = v
    return out


def check_against_results_auto(datasets: list[str], trim: str, is_real: bool) -> pd.DataFrame:

    xlsx = _read_results_auto_mae(trim)
    rows = []
    for ds in datasets:
        non_ppm = parse_approach_maes(ds, trim, is_real)
        non_ppm["nbeats"] = _nbeats_mae(ds, trim, is_real)
        for model, xlsx_row in _RESULTS_AUTO_ROW.items():
            ours, xl = non_ppm.get(model, {}), xlsx.get((xlsx_row, ds), {})
            for s in SERIES:
                a, b = ours.get(s), xl.get(s)
                if a is None or b is None or abs(a - b) <= max(1e-4, 0.001 * abs(b)):
                    continue
                rows.append({"dataset": ds, "approach": MAE_DISPLAY_NAMES.get(model, model),
                            "regime": None, "series": s, "ours": a, "results_auto": b})
        for regime in PPM_REGIMES:
            ppm = parse_ppm_maes(ds, trim, regime)
            for model in PPM_MODELS:
                xlsx_row = _RESULTS_AUTO_PPM_ROW.get((model, regime))
                if xlsx_row is None or model not in ppm:
                    continue
                a = ppm[model]
                xl = xlsx.get((xlsx_row, ds), {})
                for s in SERIES:
                    av, bv = a.get(s), xl.get(s)
                    if av is None or bv is None or abs(av - bv) <= max(1e-4, 0.001 * abs(bv)):
                        continue
                    rows.append({"dataset": ds, "approach": MAE_DISPLAY_NAMES.get(model, model),
                                "regime": regime, "series": s, "ours": av, "results_auto": bv})
    return pd.DataFrame(rows)
